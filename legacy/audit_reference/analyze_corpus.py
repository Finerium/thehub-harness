#!/usr/bin/env python3
"""Independent re-computation of every load-bearing number in the SIMPUL PRD.
Usage: python3 analyze_corpus.py --corpus "<Case 1 root>" --out out.json [--fast]
Prints a table; writes a JSON fixture with deterministic key ordering."""
import argparse, datetime, difflib, json, os, re, statistics, subprocess, sys, time
from collections import Counter, defaultdict
import openpyxl

TAG_RE = re.compile(r'[A-Z]{2}-\d{4}[A-Z]?')
MONTHS = {m.lower(): i + 1 for i, m in enumerate(
    ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
     'September', 'October', 'November', 'December'])}
DATE_RE = re.compile(r'(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})')
STOP = set("""the a an and or of to in on at for with by from is are was were be been as that this it its
into after before during not no per via than then which when while under over out up down off all any each
both has have had do does did but if so such also may can will would should must""".split())


def pdftotext(path, layout=True):
    cmd = ['pdftotext'] + (['-layout'] if layout else []) + [path, '-']
    return subprocess.run(cmd, capture_output=True, text=True).stdout


def inventory(base):
    files, junk = [], []
    for root, dirs, names in os.walk(base):
        dirs[:] = sorted(d for d in dirs if d != '__MACOSX')
        for n in sorted(names):
            if n == '.DS_Store' or n.startswith('._') or n == 'Thumbs.db':
                junk.append(os.path.join(root, n)); continue
            files.append(os.path.join(root, n))
    return files, junk


def norm(s):
    return re.sub(r'\s+', ' ', str(s).lower()).strip()


def tokens(s):
    return [t for t in re.findall(r'[a-z0-9]+(?:[-./][a-z0-9]+)*', s.lower()) if t not in STOP and len(t) > 1]


# ---------------- coverage scorers ----------------
def ratio_shipped(fl, corpus):
    """validate_fix.py::best_ratio verbatim (window = L+step, step = max(20, L//2))."""
    L = len(fl)
    if L < 20: return 0.0
    if fl in corpus: return 1.0
    step = max(20, L // 2); best = 0.0
    for s in range(0, max(1, len(corpus) - L + 1), step):
        w = corpus[s:s + L + step]
        sm = difflib.SequenceMatcher(None, fl, w)
        if sm.real_quick_ratio() < best or sm.quick_ratio() < best: continue
        r = sm.ratio()
        if r > best: best = r
        if best > 0.95: break
    return best


def ratio_window(fl, corpus, div, refine=False):
    """Window = exactly L chars (Ch.19 wording), step = max(1, L//div). refine: 1-char sweep around the best coarse hits."""
    L = len(fl)
    if L < 20: return 0.0
    if fl in corpus: return 1.0
    step = max(1, L // div); best = 0.0; hits = []
    for s in range(0, max(1, len(corpus) - L + 1), step):
        sm = difflib.SequenceMatcher(None, fl, corpus[s:s + L])
        if sm.real_quick_ratio() < best - 0.1 or sm.quick_ratio() < best - 0.1: continue
        r = sm.ratio(); hits.append((r, s))
        if r > best: best = r
    if refine and hits:
        for r, s0 in [h for h in hits if h[0] >= best - 0.1]:
            for s in range(max(0, s0 - step), min(len(corpus) - L, s0 + step) + 1):
                sm = difflib.SequenceMatcher(None, fl, corpus[s:s + L])
                if sm.real_quick_ratio() < best or sm.quick_ratio() < best: continue
                r = sm.ratio()
                if r > best: best = r
    return best


def containment(ftoks, ctoks, cset, mult=2):
    """Token containment: max over token windows (len = mult*n, step 1) of |F ∩ W| / |F|. mult=None: whole corpus."""
    F = set(ftoks)
    if not F: return 0.0
    if mult is None: return len(F & cset) / len(F)
    n = max(1, len(ftoks)) * mult
    if len(ctoks) <= n: return len(F & set(ctoks)) / len(F)
    best = 0
    win = Counter(ctoks[:n]); cur = len(F & set(win))
    best = cur
    for i in range(n, len(ctoks)):
        out, inn = ctoks[i - n], ctoks[i]
        win[out] -= 1
        if win[out] == 0: del win[out]
        win[inn] += 1
        cur = sum(1 for t in F if t in win)
        if cur > best: best = cur
        if best == len(F): break
    return best / len(F)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--corpus', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--fast', action='store_true', help='skip the slow refined char scorer')
    a = ap.parse_args()
    base = a.corpus; t0 = time.time()
    out = {}

    # 1. inventory
    files, junk = inventory(base)
    opl_files = sorted(f for f in files if os.path.basename(f).startswith('OPL-') and f.endswith('.pdf'))
    ext = Counter(os.path.splitext(f)[1].lower() for f in files)
    out['inventory'] = {'files': len(files), 'junk_skipped': sorted(os.path.relpath(j, base) for j in junk),
                        'by_ext': dict(sorted(ext.items())), 'opl_pdfs': len(opl_files),
                        'core_docs': len(files) - len(opl_files) - 2}
    print(f"FILES {len(files)} (junk skipped {len(junk)}: {[os.path.relpath(j, base) for j in junk]}) ext={dict(ext)} OPL={len(opl_files)}")

    # 2. workbook
    wb = openpyxl.load_workbook(os.path.join(base, 'Maintenance History (All Equipment).xlsx'), data_only=True)
    ws = wb['Maintenance History (All Equipm']
    hdr = [c.value for c in ws[1]]; idx = {h: i for i, h in enumerate(hdr)}
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    rows = [r for r in rows if r[idx['WO_Number']]]
    g = lambda r, c: r[idx[c]]
    def num(v):
        try: return float(str(v).replace(',', ''))
        except Exception: return None
    bd = [r for r in rows if g(r, 'Breakdown') == 'Yes']
    inc = [r for r in rows if g(r, 'Breakdown') is None]
    empty_fields = Counter()
    for r in inc:
        for h in hdr:
            if g(r, h) in (None, ''): empty_fields[h] += 1
    leads = sorted((g(r, 'Start_Date') - g(r, 'Report_Date')).total_seconds() / 3600 for r in rows)
    out['workbook'] = {
        'rows': len(rows), 'columns': len(hdr), 'breakdowns': len(bd),
        'bd_downtime_h': sum(num(g(r, 'Downtime_Hours')) or 0 for r in bd),
        'bd_cost_idr': int(sum(num(g(r, 'Total_Cost_IDR')) or 0 for r in bd)),
        'all_cost_idr': int(sum(num(g(r, 'Total_Cost_IDR')) or 0 for r in rows)),
        'labour_h': sum(num(g(r, 'Labor_Hours')) or 0 for r in rows),
        'work_types': dict(Counter(g(r, 'Work_Type') for r in rows).most_common()),
        'incomplete': len(inc), 'incomplete_by_type': dict(Counter(g(r, 'Work_Type') for r in inc).most_common()),
        'incomplete_empty_fields': dict(sorted(empty_fields.items())),
        'lead_h': {'median': statistics.median(leads), 'ge24': sum(1 for x in leads if x >= 24),
                   'ge24_pct': round(100 * sum(1 for x in leads if x >= 24) / len(leads), 1), 'min': leads[0], 'max': leads[-1]},
        'window': [str(min(g(r, 'Report_Date') for r in rows)), str(max(g(r, 'Report_Date') for r in rows))],
        'last_breakdown': str(max(g(r, 'Report_Date') for r in bd)),
    }
    print('WORKBOOK', json.dumps(out['workbook']))

    # equipment master
    master = {}
    for r in rows:
        t = g(r, 'Equipment_Tag')
        d = master.setdefault(t, {'name': g(r, 'Equipment_Name'), 'floc': g(r, 'Functional_Location'),
                                  'area': f"{g(r, 'Area_Code')} {g(r, 'Area_Name')}", 'interlocks': Counter(),
                                  'wo': 0, 'bd': 0, 'bd_h': 0.0, 'bd_cost': 0, 'all_cost': 0, 'work_types': Counter()})
        d['wo'] += 1; d['work_types'][g(r, 'Work_Type')] += 1
        if g(r, 'Related_Interlock'): d['interlocks'][str(g(r, 'Related_Interlock'))] += 1
        d['all_cost'] += int(num(g(r, 'Total_Cost_IDR')) or 0)
        if g(r, 'Breakdown') == 'Yes':
            d['bd'] += 1; d['bd_h'] += num(g(r, 'Downtime_Hours')) or 0; d['bd_cost'] += int(num(g(r, 'Total_Cost_IDR')) or 0)
    for t in master:
        master[t]['interlocks'] = dict(master[t]['interlocks'].most_common()); master[t]['work_types'] = dict(master[t]['work_types'].most_common())
    out['equipment_master'] = {t: master[t] for t in sorted(master)}
    for t in sorted(master):
        d = master[t]; print(f"MASTER {t} wo={d['wo']} bd={d['bd']} h={d['bd_h']} cost={d['bd_cost']:,} il={d['interlocks']}")

    # 3. OPL text, dates
    opl_text, opl_raw, opl_date = {}, {}, {}
    for f in opl_files:
        key = re.search(r'(OPL-[A-Z]{2}-\d{4}[A-Z]?-\d{2})', os.path.basename(f)).group(1)
        t = pdftotext(f); opl_text[key] = t; opl_raw[key] = pdftotext(f, layout=False)
        d = None; lines = t.splitlines()
        for i, l in enumerate(lines):
            if 'Date of Sharing' in l:
                for j in range(i, min(i + 4, len(lines))):
                    m = DATE_RE.search(lines[j])
                    if m and m.group(2).lower() in MONTHS:
                        d = datetime.date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1))); break
            if d: break
        opl_date[key] = d
    dated = {k: v for k, v in opl_date.items() if v}; undated = sorted(k for k, v in opl_date.items() if not v)
    ds = sorted(dated.values()); iv = sorted((b - a).days for a, b in zip(ds, ds[1:]))
    last_bd = max(g(r, 'Report_Date') for r in bd).date()
    out['opl_dates'] = {'dated': len(dated), 'undated': undated, 'min': str(ds[0]), 'max': str(ds[-1]),
                        'span_d': (ds[-1] - ds[0]).days, 'gap_after_last_breakdown_d': (ds[0] - last_bd).days,
                        'interval_median_d': statistics.median(iv), 'interval_mode': Counter(iv).most_common(2)}
    print('DATES', json.dumps(out['opl_dates']))

    # 4. latency pairs
    pairs = [('YD-2301', 'OPL-YD-2301-04', ['tube', 'crack']), ('DC-3401A', 'OPL-DC-3401A-07', ['manway']),
             ('EA-5601', 'OPL-EA-5601-05', ['tube-to-tubesheet', 'plug']), ('LV-6701', 'OPL-LV-6701-03', ['seiz']),
             ('EA-5601', 'OPL-EA-5601-06', ['foul']), ('FA-8901', 'OPL-FA-8901-07', ['manway'])]
    lat = []
    for tag, opl, kws in pairs:
        c = sorted([r for r in rows if g(r, 'Equipment_Tag') == tag and g(r, 'Breakdown') == 'Yes' and
                    any(k in (str(g(r, 'Problem_Description')) + ' ' + str(g(r, 'Root_Cause'))).lower() for k in kws)],
                   key=lambda r: g(r, 'Report_Date'))
        cands = [(g(r, 'WO_Number'), str(g(r, 'Report_Date').date()), (opl_date[opl] - g(r, 'Report_Date').date()).days) for r in c]
        lat.append({'tag': tag, 'opl': opl, 'opl_date': str(opl_date[opl]), 'chosen': cands[0], 'all_candidates': cands})
    out['latency'] = {'pairs': lat, 'median_d': statistics.median(p['chosen'][2] for p in lat)}
    for p in lat: print('LAT', p)

    # 5. per-tag corpora
    tag_full, tag_toks = defaultdict(str), defaultdict(list)
    for k in sorted(opl_text):
        tg = re.search(r'OPL-([A-Z]{2}-\d{4}[A-Z]?)-', k).group(1)
        tag_full[tg] += ' ' + norm(opl_text[k]); tag_toks[tg] += tokens(opl_text[k])
    tag_set = {t: set(v) for t, v in tag_toks.items()}
    out['corpus_size'] = {t: {'chars': len(tag_full[t]), 'tokens': len(tag_toks[t])} for t in sorted(tag_full)}

    # verbatim
    verb = [g(r, 'WO_Number') for r in rows if len(str(g(r, 'Corrective_Action') or '')) >= 20 and
            norm(g(r, 'Corrective_Action')) in tag_full[g(r, 'Equipment_Tag')]]
    out['verbatim_corrective_action'] = {'count': len(verb), 'wos': verb}
    print('VERBATIM', len(verb))

    # 6. coverage under every scorer
    methods = {'shipped_L+step': lambda fl, c, tg: ratio_shipped(fl, c),
               'winL_step2': lambda fl, c, tg: ratio_window(fl, c, 2),
               'winL_step4': lambda fl, c, tg: ratio_window(fl, c, 4),
               'tok_contain_w2': lambda fl, c, tg: containment(tokens(fl), tag_toks[tg], tag_set[tg], 2),
               'tok_contain_global': lambda fl, c, tg: containment(tokens(fl), tag_toks[tg], tag_set[tg], None)}
    if not a.fast: methods['winL_refined'] = lambda fl, c, tg: ratio_window(fl, c, 4, refine=True)
    scores = {m: {} for m in methods}
    for m, fn in methods.items():
        t1 = time.time()
        for r in rows:
            tg = g(r, 'Equipment_Tag'); c = tag_full[tg]; b = 0.0
            for cf in ('Problem_Description', 'Root_Cause', 'Corrective_Action'):
                v = norm(g(r, cf) or '')
                if len(v) >= 20: b = max(b, fn(v, c, tg))
                if b >= 1.0: break
            scores[m][g(r, 'WO_Number')] = round(b, 4)
        print(f"scored {m} in {time.time() - t1:.1f}s")
    out['scores'] = scores
    table = {}
    for m in methods:
        table[m] = {}
        for th in (0.55, 0.62, 0.70, 0.75):
            unc = [r for r in rows if scores[m][g(r, 'WO_Number')] <= th]
            ubd = [r for r in bd if scores[m][g(r, 'WO_Number')] <= th]
            table[m][str(th)] = {'uncovered': len(unc), 'pct': round(100 * len(unc) / len(rows), 1),
                                 'unc_breakdowns': len(ubd), 'unc_bd_h': sum(num(g(r, 'Downtime_Hours')) or 0 for r in ubd),
                                 'unc_bd_idr': int(sum(num(g(r, 'Total_Cost_IDR')) or 0 for r in ubd)),
                                 'unc_by_tag': dict(Counter(g(r, 'Equipment_Tag') for r in unc).most_common()),
                                 'unc_bd_wos': sorted(g(r, 'WO_Number') for r in ubd)}
            x = table[m][str(th)]
            print(f"COV {m:20s} t={th}: unc={x['uncovered']}/211 ({x['pct']}%) bd={x['unc_breakdowns']}/31 h={x['unc_bd_h']} IDR={x['unc_bd_idr']:,} by_tag={x['unc_by_tag']}")
    out['coverage'] = table

    # 7. cross-reference audit (robust)
    xref = {}
    for k in sorted(opl_text):
        subj = re.search(r'OPL-([A-Z]{2}-\d{4}[A-Z]?)-', k).group(1)
        t = re.sub(r'\s+', ' ', opl_text[k])
        strict = re.search(r'Cross-ref tag\s+([A-Z]{2}-\d{4}[A-Z]?)', opl_text[k])
        # robust: every "<TAG> across" mention, any whitespace, plus any tag next to 'Cross-ref'/'DUMMY tag'
        mentions = set(re.findall(r'([A-Z]{2}-\d{4}[A-Z]?)\s*across', t))
        mentions |= set(re.findall(r'(?:Cross-ref|DUMMY)\s*tag\s*([A-Z]{2}-\d{4}[A-Z]?)', t))
        mentions |= set(re.findall(r'Cross-ref\s+([A-Z]{2}-\d{4}[A-Z]?)', t))
        xref[k] = {'subject': subj, 'strict_first': strict.group(1) if strict else None,
                   'mentions': sorted(mentions), 'wrong': sorted(m for m in mentions if m != subj),
                   'has_footer': bool(mentions), 'has_dummy': 'DUMMY' in opl_text[k]}
    per_set_wrong = Counter(v['subject'] for v in xref.values() if v['wrong'])
    strict_wrong = Counter(v['subject'] for v in xref.values() if v['strict_first'] and v['strict_first'] != v['subject'])
    out['xref'] = {'per_doc': xref, 'robust_wrong_per_set': dict(sorted(per_set_wrong.items())), 'robust_wrong_total': sum(per_set_wrong.values()),
                   'strict_wrong_per_set': dict(sorted(strict_wrong.items())), 'strict_wrong_total': sum(strict_wrong.values()),
                   'no_footer': sorted(k for k, v in xref.items() if not v['has_footer']),
                   'strict_unparsed': sorted(k for k, v in xref.items() if not v['strict_first']),
                   'dummy_docs': sorted(k for k, v in xref.items() if v['has_dummy']),
                   'both_right_and_wrong': sorted(k for k, v in xref.items() if v['wrong'] and v['subject'] in v['mentions'])}
    print('XREF robust', out['xref']['robust_wrong_per_set'], out['xref']['robust_wrong_total'], 'strict', out['xref']['strict_wrong_per_set'],
          'no_footer', out['xref']['no_footer'], 'strict_unparsed', out['xref']['strict_unparsed'], 'dummy', out['xref']['dummy_docs'],
          'both', out['xref']['both_right_and_wrong'])

    out['elapsed_s'] = round(time.time() - t0, 1)
    with open(a.out, 'w') as fh: json.dump(out, fh, indent=1, sort_keys=False, default=str)
    print('wrote', a.out, out['elapsed_s'], 's')


if __name__ == '__main__':
    main()
