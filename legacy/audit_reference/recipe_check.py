#!/usr/bin/env python3
"""A: score histograms + fine thresholds from out.json. B: per-document token containment (frozen-recipe candidate).
C: order/whitespace sensitivity of the shipped best_ratio."""
import json, os, re, glob, random, statistics, difflib
from collections import Counter, defaultdict
import openpyxl
from analyze_corpus import ratio_shipped, norm, tokens, STOP

S = "/private/tmp/claude-501/-Users-ghaisan-Documents-ChandraAsri-Competition/f63551eb-0fe4-4530-8e18-97ee139ea91d/scratchpad"
BASE = "/Users/ghaisan/Documents/ChandraAsri-Competition/CALIBER 2026 - The Case/Supporting Data/Case 1_ Manufacturing Knowledge Hub"
out = json.load(open('out.json'))
wb = openpyxl.load_workbook(os.path.join(BASE, "Maintenance History (All Equipment).xlsx"), data_only=True)
ws = wb['Maintenance History (All Equipm']; hdr = [c.value for c in ws[1]]; idx = {h: i for i, h in enumerate(hdr)}
rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
g = lambda r, c: r[idx[c]]
num = lambda v: float(str(v).replace(',', '')) if v not in (None, '') else 0.0
bd = [r for r in rows if g(r, 'Breakdown') == 'Yes']
THS = (0.5, 0.55, 0.6, 0.62, 0.65, 0.7, 0.75, 0.8)

def summarize(scores, ths=THS, label=''):
    res = {}
    for th in ths:
        unc = [r for r in rows if scores[g(r, 'WO_Number')] <= th]; ubd = [r for r in bd if scores[g(r, 'WO_Number')] <= th]
        res[th] = dict(unc=len(unc), pct=round(100 * len(unc) / 211, 1), bd=len(ubd), h=sum(num(g(r, 'Downtime_Hours')) for r in ubd),
                       idr=int(sum(num(g(r, 'Total_Cost_IDR')) for r in ubd)), by_tag=dict(sorted(Counter(g(r, 'Equipment_Tag') for r in unc).items())),
                       bd_wos=sorted(g(r, 'WO_Number') for r in ubd))
        print(f"  {label:14s} t={th:<5} unc={res[th]['unc']:3d}/211 ({res[th]['pct']:4.1f}%) bd={res[th]['bd']:2d}/31 h={res[th]['h']:6.1f} IDR={res[th]['idr']:>12,} by_tag={res[th]['by_tag']}")
    return res

# ---------- A: histograms ----------
print("A. score histograms (bins of 0.05; count of WOs whose best score falls in bin) and exact-1.0 count")
for m, sc in out['scores'].items():
    h = Counter(min(int(v * 20), 19) / 20 for v in sc.values())
    print(f"  {m:20s} 1.0:{sum(1 for v in sc.values() if v >= 1.0):3d} | " + ' '.join(f"{b:.2f}:{h.get(b,0)}" for b in sorted(h)))
print("A2. fine thresholds for the two candidates")
summarize(out['scores']['winL_refined'], label='winL_refined')
summarize(out['scores']['tok_contain_w2'], label='tok_w2(concat)')

# ---------- B: per-document token containment ----------
opl_docs = defaultdict(list)  # tag -> [(key, tokens)]
for f in sorted(glob.glob(S + "/case1txt/*OPL-*.txt")):
    k = re.search(r'(OPL-[A-Z]{2}-\d{4}[A-Z]?-\d{2})', os.path.basename(f)).group(1)
    tg = re.search(r'OPL-([A-Z]{2}-\d{4}[A-Z]?)-', k).group(1)
    opl_docs[tg].append((k, tokens(open(f).read())))

def contain_doc(ftoks, dtoks, mult=2):
    F = set(ftoks)
    if not F: return 0.0
    n = len(ftoks) * mult
    if len(dtoks) <= n: return len(F & set(dtoks)) / len(F)
    best = 0; win = Counter(dtoks[:n]); best = sum(1 for t in F if t in win)
    for i in range(n, len(dtoks)):
        o, inn = dtoks[i - n], dtoks[i]
        win[o] -= 1
        if not win[o]: del win[o]
        win[inn] += 1
        c = sum(1 for t in F if t in win)
        if c > best: best = c
        if best == len(F): break
    return best / len(F)

def score_all(mult=2, drop_tag=False, order=None):
    sc = {}; unit = {}
    for r in rows:
        tg = g(r, 'Equipment_Tag'); b = 0.0; bu = None
        docs = opl_docs[tg] if order is None else order(opl_docs[tg])
        for cf in ('Problem_Description', 'Root_Cause', 'Corrective_Action'):
            v = str(g(r, cf) or '')
            ft = tokens(v)
            if drop_tag: ft = [t for t in ft if t != tg.lower()]
            if len(ft) < 3: continue
            for k, dt in docs:
                c = contain_doc(ft, dt, mult)
                if c > b: b, bu = c, (cf, k)
        sc[g(r, 'WO_Number')] = round(b, 4); unit[g(r, 'WO_Number')] = bu
    return sc, unit

print("B. per-document token containment (window = 2n tokens, step 1, stopwords removed, tokens>=2 chars)")
sc2, unit2 = score_all(2)
resB = summarize(sc2, label='tok_doc_w2')
print("  determinism: shuffled doc order gives identical scores:", score_all(2, order=lambda d: random.sample(d, len(d)))[0] == sc2)
print("  exact 1.0 count:", sum(1 for v in sc2.values() if v >= 1.0), "; verbatim WOs all 1.0:", all(sc2[w] >= 1.0 for w in out['verbatim_corrective_action']['wos']))
print("B2. variants: window 3n; drop subject tag token")
summarize(score_all(3)[0], ths=(0.55, 0.62, 0.7, 0.75), label='tok_doc_w3')
summarize(score_all(2, drop_tag=True)[0], ths=(0.55, 0.62, 0.7, 0.75), label='tok_doc_w2_notag')
# which breakdowns differ between shipped(t=.62) and recipe(t=.62)
ship = set(out['coverage']['shipped_L+step']['0.62']['unc_bd_wos']); rec = set(resB[0.62]['bd_wos'])
print("  breakdowns uncovered shipped-only:", sorted(ship - rec), " recipe-only:", sorted(rec - ship))
for w in sorted(ship - rec):
    r = next(r for r in rows if g(r, 'WO_Number') == w)
    print(f"    {w} {g(r,'Equipment_Tag')} score={sc2[w]} unit={unit2[w]} | {str(g(r,'Problem_Description'))[:70]}")
json.dump({'tok_doc_w2_scores': sc2, 'tok_doc_w2_units': {k: list(v) if v else None for k, v in unit2.items()}, 'tok_doc_w2_table': {str(k): v for k, v in resB.items()}}, open('recipe.json', 'w'), indent=1)

# ---------- C: order / whitespace sensitivity of the shipped best_ratio ----------
print("C. shipped best_ratio: headline at t=0.62 under different (equally valid) corpus concatenations")
texts = {}
for f in glob.glob(S + "/case1txt/*OPL-*.txt"):
    k = re.search(r'(OPL-[A-Z]{2}-\d{4}[A-Z]?-\d{2})', os.path.basename(f)).group(1); texts[k] = open(f).read()
def build(order, strip):
    tf = defaultdict(str)
    for k in order:
        tg = re.search(r'OPL-([A-Z]{2}-\d{4}[A-Z]?)-', k).group(1)
        t = re.sub(r'\s+', ' ', texts[k].lower()); tf[tg] += ' ' + (t.strip() if strip else t)
    return tf
def run(tf):
    sc = {}
    for r in rows:
        c = tf[g(r, 'Equipment_Tag')]; b = 0.0
        for cf in ('Problem_Description', 'Root_Cause', 'Corrective_Action'):
            v = str(g(r, cf) or '')
            if len(v) >= 20: b = max(b, ratio_shipped(norm(v), c))
            if b > 0.95: break
        sc[g(r, 'WO_Number')] = b
    return sc
keys = sorted(texts)
for name, order, strip in [('sorted,nostrip', keys, False), ('sorted,strip', keys, True), ('reversed,nostrip', keys[::-1], False), ('shuffled(seed1),nostrip', random.Random(1).sample(keys, len(keys)), False)]:
    sc = run(build(order, strip))
    print(f"  {name:26s}: " + ' '.join(f"t={th}:{sum(1 for v in sc.values() if v <= th)}" for th in (0.55, 0.62, 0.7, 0.75)) + f"  bd@.62={sum(1 for r in bd if sc[g(r,'WO_Number')] <= 0.62)}")
