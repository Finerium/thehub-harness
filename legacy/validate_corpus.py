import os, re, subprocess, statistics, datetime, difflib, json
from collections import Counter, defaultdict
import openpyxl

BASE = "/home/claude/caliber_case/CALIBER 2026 - The Case/Supporting Data/Case 1_ Manufacturing Knowledge Hub"

# ---------- 1. File inventory ----------
all_files = []
for root, dirs, files in os.walk(BASE):
    if "__MACOSX" in root: continue
    for f in files:
        if f.startswith("._"): continue
        all_files.append(os.path.join(root, f))
opl_files = [f for f in all_files if re.search(r'OPL-', os.path.basename(f), re.I) and f.endswith('.pdf')]
print(f"FILES: total={len(all_files)}  OPL_pdfs={len(opl_files)}")
sets = sorted([d for d in os.listdir(BASE) if d.startswith("Set_")])
for s in sets:
    core = [f for f in os.listdir(os.path.join(BASE,s)) if os.path.isfile(os.path.join(BASE,s,f)) and not f.startswith("._")]
    ga = [f for f in core if 'Drawing' in f]
    print(f"  {s}: core={len(core)} drawing_files={ga}")

# ---------- 2. Excel ----------
wb = openpyxl.load_workbook(os.path.join(BASE, "Maintenance History (All Equipment).xlsx"), data_only=True)
ws = wb['Maintenance History (All Equipm']
hdr = [c.value for c in ws[1]]
idx = {h:i for i,h in enumerate(hdr)}
rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
print(f"\nWO rows: {len(rows)}")

def g(r, col): return r[idx[col]]
def num(v):
    if v is None: return None
    try: return float(str(v).replace(',',''))
    except: return None

breakdowns = [r for r in rows if g(r,'Breakdown')=='Yes']
incomplete = [r for r in rows if g(r,'Breakdown') is None]
bd_dt = sum(num(g(r,'Downtime_Hours')) or 0 for r in breakdowns)
bd_cost = sum(num(g(r,'Total_Cost_IDR')) or 0 for r in breakdowns)
all_cost = sum(num(g(r,'Total_Cost_IDR')) or 0 for r in rows)
print(f"Breakdowns: {len(breakdowns)}  downtime_h={bd_dt}  bd_cost_IDR={bd_cost:,.0f}  all_cost_IDR={all_cost:,.0f}")
print(f"Incomplete rows: {len(incomplete)}  by worktype={Counter(g(r,'Work_Type') for r in incomplete)}")

leads = []
for r in rows:
    rp, st = g(r,'Report_Date'), g(r,'Start_Date')
    if isinstance(rp,datetime.datetime) and isinstance(st,datetime.datetime):
        leads.append((st-rp).total_seconds()/3600)
leads.sort()
print(f"Lead time: n={len(leads)} median={statistics.median(leads)}h  >=24h={sum(1 for x in leads if x>=24)} ({100*sum(1 for x in leads if x>=24)/len(leads):.1f}%)  min={min(leads)} max={max(leads)}")

last_bd = max(g(r,'Report_Date') for r in breakdowns)
print(f"WO window: {min(g(r,'Report_Date') for r in rows)} -> {max(g(r,'Report_Date') for r in rows)}; last breakdown report: {last_bd}")

# YD-2301 interlock + EA-5601 interlock
print("YD-2301 Related_Interlock:", Counter(g(r,'Related_Interlock') for r in rows if g(r,'Equipment_Tag')=='YD-2301'))
print("EA-5601 Related_Interlock:", Counter(g(r,'Related_Interlock') for r in rows if g(r,'Equipment_Tag')=='EA-5601'))

# GA-1201A causal chain
for won in ['WO-240002','WO-240003','WO-240004']:
    for r in rows:
        if g(r,'WO_Number')==won:
            print(f"{won} [{g(r,'Report_Date')}] tag={g(r,'Equipment_Tag')} prob={g(r,'Problem_Description')} | RC={g(r,'Root_Cause')}")

# DC-3401A PSV gap
psv = [r for r in rows if g(r,'Equipment_Tag')=='DC-3401A' and 'PSV' in str(g(r,'Problem_Description'))+str(g(r,'Root_Cause'))]
print(f"DC-3401A PSV-related WOs: {[(g(r,'WO_Number'),g(r,'Breakdown'),str(g(r,'Report_Date'))[:10],g(r,'Problem_Description')) for r in psv]}")
set3_opls = [os.path.basename(f) for f in opl_files if 'DC-3401A' in f]
print(f"DC-3401A OPL titles: {set3_opls}")

# ---------- 3. OPL text + Date of Sharing ----------
def pdf_text(p):
    return subprocess.run(['pdftotext','-layout',p,'-'],capture_output=True,text=True).stdout

opl_text, opl_date = {}, {}
date_pat = re.compile(r'(\d{1,2}\s+[A-Za-z]+\s+\d{4}|\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{4})')
months = {m.lower():i+1 for i,m in enumerate(['January','February','March','April','May','June','July','August','September','October','November','December'])}
def parse_date(s):
    s=s.strip()
    m=re.match(r'(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})',s)
    if m and m.group(2).lower() in months: return datetime.date(int(m.group(3)),months[m.group(2).lower()],int(m.group(1)))
    m=re.match(r'(\d{4})-(\d{2})-(\d{2})',s)
    if m: return datetime.date(int(m.group(1)),int(m.group(2)),int(m.group(3)))
    m=re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})',s)
    if m: return datetime.date(int(m.group(3)),int(m.group(2)),int(m.group(1)))
    return None

samples=0
for f in opl_files:
    name = os.path.basename(f)
    key = re.search(r'(OPL-[A-Z]{2}-\d{4}[A-Z]?-\d{2})', name)
    key = key.group(1) if key else name
    t = pdf_text(f)
    opl_text[key] = t
    d=None
    for line in t.splitlines():
        if 'sharing' in line.lower():
            if samples<3: print("FOOTER SAMPLE:", line.strip()[:120]); samples+=1
            m=date_pat.search(line)
            if m: d=parse_date(m.group(1))
    if d is None:
        # search anywhere near 'Sharing' across text
        m=re.search(r'Sharing[^\n]{0,60}', t)
    opl_date[key]=d

dated = {k:v for k,v in opl_date.items() if v}
undated = [k for k,v in opl_date.items() if not v]
print(f"\nOPL total={len(opl_text)}  dated={len(dated)}  undated={undated}")
if dated:
    dts=sorted(dated.values())
    print(f"Date of Sharing: min={dts[0]}  max={dts[-1]}  span_days={(dts[-1]-dts[0]).days}")
    print(f"Earliest OPL vs last breakdown gap: {(dts[0]-last_bd.date()).days} days")

# ---------- 4. Six latency pairs ----------
pairs = [('YD-2301','OPL-YD-2301-04',['tube','crack']),
         ('DC-3401A','OPL-DC-3401A-07',['manway']),
         ('EA-5601','OPL-EA-5601-05',['tube']),
         ('LV-6701','OPL-LV-6701-03',['seiz','packing']),
         ('EA-5601','OPL-EA-5601-06',['foul']),
         ('FA-8901','OPL-FA-8901-07',['manway'])]
lat=[]
print("\nLATENCY PAIRS:")
for tag,opl,kws in pairs:
    cands=[r for r in rows if g(r,'Equipment_Tag')==tag and g(r,'Breakdown')=='Yes' and any(k in (str(g(r,'Problem_Description'))+' '+str(g(r,'Root_Cause'))).lower() for k in kws)]
    cands.sort(key=lambda r:g(r,'Report_Date'))
    if cands and opl_date.get(opl):
        fd=g(cands[0],'Report_Date').date(); od=opl_date[opl]; L=(od-fd).days; lat.append(L)
        print(f"  {tag} {str(fd)} -> {opl} {od}  latency={L}d  ({g(cands[0],'Problem_Description')[:50]})")
    else:
        print(f"  {tag}/{opl}: MATCH FAILED (cands={len(cands)}, opl_date={opl_date.get(opl)})")
if lat: print(f"  median latency = {statistics.median(sorted(lat))} days over {len(lat)} pairs")

# ---------- 5. Coverage gap ----------
tag_lines={}
tag_full={}
for k,t in opl_text.items():
    m=re.search(r'OPL-([A-Z]{2}-\d{4}[A-Z]?)-',k)
    tg=m.group(1) if m else '?'
    ls=[l.strip() for l in t.splitlines() if len(l.strip())>=15]
    tag_lines.setdefault(tg,[]).extend(ls)
    tag_full[tg]=tag_full.get(tg,'')+' '+re.sub(r'\s+',' ',t.lower())

def covered(r, th):
    tg=g(r,'Equipment_Tag')
    fields=[str(g(r,c) or '') for c in ('Problem_Description','Root_Cause','Corrective_Action')]
    fields=[f for f in fields if len(f)>=15]
    for f in fields:
        fl=f.lower()
        for line in tag_lines.get(tg,[]):
            ll=line.lower()
            sm=difflib.SequenceMatcher(None, fl, ll)
            if sm.real_quick_ratio()<=th: continue
            if sm.quick_ratio()<=th: continue
            if sm.ratio()>th: return True
    return False

for th in (0.55,0.62,0.70,0.75):
    cov=[r for r in rows if covered(r,th)]
    unc=len(rows)-len(cov)
    unc_bd=[r for r in breakdowns if not covered(r,th)]
    print(f"\nTHRESHOLD {th}: covered={len(cov)}  uncovered={unc} ({100*unc/len(rows):.1f}%)  uncovered_breakdowns={len(unc_bd)}/{len(breakdowns)}  unc_bd_downtime={sum(num(g(r,'Downtime_Hours')) or 0 for r in unc_bd)}h  unc_bd_cost={sum(num(g(r,'Total_Cost_IDR')) or 0 for r in unc_bd):,.0f}")

# ---------- 6. Verbatim mining ----------
verb=0; opls_with=set()
for r in rows:
    ca=str(g(r,'Corrective_Action') or '')
    if len(ca)<20: continue
    tg=g(r,'Equipment_Tag')
    if re.sub(r'\s+',' ',ca.lower()) in tag_full.get(tg,''):
        verb+=1
print(f"\nWOs whose Corrective_Action appears VERBATIM in same-tag OPL text: {verb}/{len(rows)}")
