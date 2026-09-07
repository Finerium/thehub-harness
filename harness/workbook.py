"""Maintenance History workbook: typed rows, closeout flags and the plan's populations (Revision Plan section 5.4).

Populations (frozen): failure = Work_Type in {Corrective, Overhaul} or Breakdown == Yes; unplanned_failure = failure minus
rows whose Problem_Description starts with Scheduled|Statutory|Turnaround|Grid inspection; planned_flagged = Breakdown == Yes
and planned; unplanned_breakdowns = Breakdown == Yes and not planned.
"""
import re

import openpyxl

from .config import WORKBOOK

NARR = ("Problem_Description", "Root_Cause", "Corrective_Action")
OUTCOME_FIELDS = (
    "Breakdown", "Downtime_Hours", "Labor_Hours", "Labor_Cost_IDR", "Material_Cost_IDR", "Total_Cost_IDR",
    "Reported_By", "Executed_By", "Approved_By", "Related_Interlock", "Remarks",
)
PLANNED = re.compile(r"^(Scheduled|Statutory|Turnaround|Grid inspection)", re.IGNORECASE)


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load(path=WORKBOOK):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r, strict=True))
        d["Downtime_Hours"] = _num(d.get("Downtime_Hours"))
        d["Labor_Hours"] = _num(d.get("Labor_Hours"))
        for c in ("Labor_Cost_IDR", "Material_Cost_IDR", "Total_Cost_IDR"):
            v = _num(d.get(c))
            d[c] = int(v) if v is not None else None
        d["closeout_complete"] = all(d.get(c) not in (None, "") for c in OUTCOME_FIELDS)
        d["is_planned"] = bool(PLANNED.match(str(d.get("Problem_Description") or "")))
        d["is_failure"] = d.get("Work_Type") in ("Corrective", "Overhaul") or d.get("Breakdown") == "Yes"
        if d.get("Breakdown") == "Yes":
            d["breakdown_kind"] = "planned_flagged" if d["is_planned"] else "unplanned"
        else:
            d["breakdown_kind"] = None
        rows.append(d)
    return rows


def explanation(path=WORKBOOK):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Explanation"]
    return [(r[0], r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def populations(rows):
    fail = [w for w in rows if w["is_failure"]]
    return {
        "all": rows,
        "failure": fail,
        "unplanned_failure": [w for w in fail if not w["is_planned"]],
        "planned_flagged": [w for w in rows if w["breakdown_kind"] == "planned_flagged"],
        "unplanned_breakdowns": [w for w in rows if w["breakdown_kind"] == "unplanned"],
    }


def ids(rows):
    return sorted(w["WO_Number"] for w in rows)
